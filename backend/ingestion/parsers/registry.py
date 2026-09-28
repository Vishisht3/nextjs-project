"""
ingestion/parsers/registry.py
Registry mapping file extensions to document parsers.
"""
from __future__ import annotations

import json
from pathlib import Path
from typing import Dict, Type

from ingestion.parsers.base import BaseParser, ParsedDocument, UnsupportedFormatError


class PlainTextParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return ParsedDocument(
            text=content,
            source=path.name,
            metadata={"filepath": str(path), "extension": path.suffix},
        )


class MarkdownParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        return ParsedDocument(
            text=content,
            source=path.name,
            metadata={"filepath": str(path), "extension": path.suffix, "format": "markdown"},
        )


class JsonParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            data = json.load(f)
        if isinstance(data, dict):
            text = data.get("text") or data.get("content") or json.dumps(data, indent=2)
            meta = {k: v for k, v in data.items() if k not in {"text", "content"}}
        elif isinstance(data, list):
            text = "\n\n".join(str(item) for item in data)
            meta = {"items_count": len(data)}
        else:
            text = str(data)
            meta = {}
        return ParsedDocument(
            text=text,
            source=path.name,
            metadata={"filepath": str(path), "extension": path.suffix, **meta},
        )


class PdfParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        try:
            import pypdf
            reader = pypdf.PdfReader(str(path))
            pages = [page.extract_text() or "" for page in reader.pages]
            full_text = "\n\n".join(pages)
            return ParsedDocument(
                text=full_text,
                source=path.name,
                metadata={"filepath": str(path), "pages": len(reader.pages)},
            )
        except ImportError:
            # Fallback if pypdf is not installed
            with open(path, "rb") as f:
                raw = f.read().decode("latin1", errors="ignore")
            return ParsedDocument(
                text=raw,
                source=path.name,
                metadata={"filepath": str(path), "note": "pypdf not installed, raw read"},
            )


class DocxParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError("DOCX uploads require python-docx") from exc

        document = Document(str(path))
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_rows = [" | ".join(cell.text.strip() for cell in row.cells) for table in document.tables for row in table.rows]
        text = "\n".join(paragraphs + table_rows)
        return ParsedDocument(
            text=text,
            source=path.name,
            metadata={"filepath": str(path), "extension": path.suffix, "format": "docx"},
        )


class PptxParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise ImportError("PPTX uploads require python-pptx") from exc

        presentation = Presentation(str(path))
        slides = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            parts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
            if parts:
                slides.append(f"[Slide {slide_number}]\n" + "\n".join(parts))
        return ParsedDocument(
            text="\n\n".join(slides),
            source=path.name,
            metadata={"filepath": str(path), "extension": path.suffix, "format": "pptx", "slides": len(presentation.slides)},
        )


class TexParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        text = content if self.keep_math else _strip_latex_comments(content)
        return ParsedDocument(
            text=text,
            source=path.name,
            metadata={"filepath": str(path), "extension": path.suffix, "format": "tex", "keep_math": self.keep_math},
        )


def _strip_latex_comments(content: str) -> str:
    return "\n".join(line.split("%", 1)[0] for line in content.splitlines())


PARSER_REGISTRY: Dict[str, Type[BaseParser]] = {
    ".txt": PlainTextParser,
    ".md": MarkdownParser,
    ".markdown": MarkdownParser,
    ".json": JsonParser,
    ".csv": PlainTextParser,
    ".pdf": PdfParser,
    ".docx": DocxParser,
    ".pptx": PptxParser,
    ".tex": TexParser,
    ".py": PlainTextParser,
    ".yaml": PlainTextParser,
    ".yml": PlainTextParser,
}


def get_parser_for(path: Path, keep_math: bool = False) -> BaseParser:
    ext = path.suffix.lower()
    parser_cls = PARSER_REGISTRY.get(ext)
    if parser_cls is None:
        raise UnsupportedFormatError(f"Unsupported file format '{ext}' for file: {path}")
    return parser_cls(keep_math=keep_math)
